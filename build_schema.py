#!/usr/bin/env python3
"""
build_schema.py — reads structure.xlsx and emits form-schema.json.

Usage:
    python build_schema.py structure.xlsx [output.json]

The HTML/JS form engine fetches the output JSON at page load and renders
itself from it. This script never needs to change for ordinary revisions —
only structure.xlsx does. Edit the spreadsheet, re-run this, commit both.

All problems found are reported together (not just the first one), and
nothing is written unless the sheet is fully clean.

KIND AXES: there is no longer a single "kind" enum. Any number of ordinary
radio/select fields can act as a "kind axis" — their option values become
valid entries for every applies_to_kinds / required_for_kinds column in the
sheet. Which fields are axes is declared in the KindAxes sheet (one field_id
per row). The active "kind set" at runtime is the union of the CURRENT
values of every axis field, not a single selector.
"""
import sys
import re
import json
from pathlib import Path

import openpyxl

SHOW_IF_EQ = re.compile(r"^\s*([A-Za-z0-9_]+)\s*(==|!=)\s*'([^']*)'\s*$")
SHOW_IF_SET = re.compile(r"^\s*([A-Za-z0-9_]+)\s+(all|not_all|in|not_in)\s*\[\s*([A-Za-z0-9_,\s]*)\s*\]\s*$")

VALID_TYPES = {"text", "email", "tel", "textarea", "date", "radio", "select", "checkbox_group", "info"}
VALID_REQUIRED = {"TRUE", "FALSE", "IF_VISIBLE"}
OPTION_TYPES = {"radio", "select", "checkbox_group"}
VALID_UI_STYLES = {"", "toggle"}


class SpecError(Exception):
    pass


def sheet_rows(ws):
    """Yield (row_number, {header: value}) for every non-empty data row."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    for i, row in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        record = {headers[c]: row[c] for c in range(len(headers)) if headers[c]}
        yield i, record


def norm_bool(v, default="FALSE"):
    if v is None or str(v).strip() == "":
        return default
    return str(v).strip().upper()


def norm_list(v):
    """'all' | 'a,b,c' -> list, stripped; 'all' passes through literally."""
    if v is None or str(v).strip() == "":
        return ["all"]
    v = str(v).strip()
    if v.lower() == "all":
        return ["all"]
    return [x.strip() for x in v.split(",") if x.strip()]


def parse_options(v):
    """'val:Label|val:Label' -> [{'value':..,'label':..}, ...]"""
    if v is None or str(v).strip() == "":
        return []
    out = []
    for part in str(v).split("|"):
        part = part.strip()
        if not part:
            continue
        if ":" not in part:
            raise SpecError(f"option '{part}' missing ':' between value and label")
        val, label = part.split(":", 1)
        out.append({"value": val.strip(), "label": label.strip()})
    return out


def resolve_show_if(raw, fields, errors, where, self_id=None, self_group_id="__not_a_field__",
                     require_groupless_ref=False):
    """
    Parse a show_if string against the fully-built `fields` dict and validate it.
    where: a label like "Fields row 12 (q5)" used in error messages.
    self_id: the id of the field/group this show_if belongs to (to catch self-reference).
    self_group_id: the group_id of the thing this show_if belongs to; pass the sentinel
        default for callers (like RepeatGroups) that have no group scope of their own.
    require_groupless_ref: if True, the referenced field must be ungrouped (used for
        RepeatGroups, since a section-level group can't unambiguously depend on one
        instance of another repeating group).
    """
    if not raw:
        return None

    m_eq = SHOW_IF_EQ.match(raw)
    m_set = SHOW_IF_SET.match(raw)
    result = None
    ref = None
    if m_eq:
        ref, op, val = m_eq.groups()
        if ref not in fields:
            errors.append(f"{where}: show_if references unknown field '{ref}'")
        else:
            ref_field = fields[ref]
            if ref_field["type"] in OPTION_TYPES:
                valid_vals = {o["value"] for o in ref_field["options"]}
                if val not in valid_vals:
                    errors.append(
                        f"{where}: show_if value '{val}' is not an option "
                        f"of '{ref}' (valid: {sorted(valid_vals)})"
                    )
        result = {"ref": ref, "op": op, "value": val}
    elif m_set:
        ref, op, raw_vals = m_set.groups()
        vals = [v.strip() for v in raw_vals.split(",") if v.strip()]
        if ref not in fields:
            errors.append(f"{where}: show_if references unknown field '{ref}'")
        else:
            ref_field = fields[ref]
            needs_type = {"all", "not_all"}
            wants_checkbox_group = op in needs_type
            if wants_checkbox_group and ref_field["type"] != "checkbox_group":
                errors.append(
                    f"{where}: show_if uses {op} on '{ref}', "
                    f"which is type '{ref_field['type']}', not checkbox_group"
                )
            elif not wants_checkbox_group and ref_field["type"] not in ("radio", "select"):
                errors.append(
                    f"{where}: show_if uses {op} on '{ref}', "
                    f"which is type '{ref_field['type']}', not radio/select"
                )
            else:
                valid_vals = {o["value"] for o in ref_field["options"]}
                bad = [v for v in vals if v not in valid_vals]
                if bad:
                    errors.append(
                        f"{where}: show_if values {bad} not valid options "
                        f"of '{ref}' (valid: {sorted(valid_vals)})"
                    )
        result = {"ref": ref, "op": op, "values": vals}
    else:
        errors.append(f"{where}: could not parse show_if '{raw}'")
        return None

    if ref == self_id:
        errors.append(f"{where}: show_if cannot reference itself")
    elif ref in fields:
        ref_group_id = fields[ref]["group_id"]
        if require_groupless_ref:
            if ref_group_id:
                errors.append(
                    f"{where}: show_if references '{ref}', which is inside repeat group "
                    f"'{ref_group_id}' — a group's own visibility can only depend on an "
                    f"ungrouped field (which instance would it mean otherwise?)"
                )
        elif self_group_id != ref_group_id:
            errors.append(
                f"{where}: show_if references '{ref}' in a different scope "
                f"(group '{ref_group_id}' vs '{self_group_id}') — a field can only show_if on "
                f"another field in the SAME repeat group, or another ungrouped field if it is "
                f"itself ungrouped"
            )
    return result


def build(xlsx_path: Path):
    errors = []
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    for required_sheet in ("KindAxes", "Sections", "Fields"):
        if required_sheet not in wb.sheetnames:
            raise SpecError(f"missing required sheet '{required_sheet}'")

    # ---- Sections (raw: collect applies_to_kinds as a plain list, validate later) ----
    sections = {}
    for rownum, r in sheet_rows(wb["Sections"]):
        sid = str(r.get("section_id", "")).strip()
        if not sid:
            errors.append(f"Sections row {rownum}: empty section_id")
            continue
        if sid in sections:
            errors.append(f"Sections row {rownum}: duplicate section_id '{sid}'")
        applies = norm_list(r.get("applies_to_kinds"))
        try:
            order = float(r.get("order", 0) or 0)
        except (TypeError, ValueError):
            errors.append(f"Sections row {rownum} ({sid}): 'order' is not a number")
            order = 0
        pdf_order_raw = r.get("pdf_order")
        try:
            pdf_order = float(pdf_order_raw) if pdf_order_raw not in (None, "") else order
        except (TypeError, ValueError):
            errors.append(f"Sections row {rownum} ({sid}): 'pdf_order' is not a number")
            pdf_order = order
        sections[sid] = {
            "id": sid,
            "order": order,
            "pdf_order": pdf_order,
            "title": str(r.get("title", "")).strip(),
            "applies_to_kinds": applies,
            "_row": rownum,
            "fields": [],
        }

    # ---- RepeatGroups (raw) ----
    groups = {}
    if "RepeatGroups" in wb.sheetnames:
        for rownum, r in sheet_rows(wb["RepeatGroups"]):
            gid = str(r.get("group_id", "")).strip()
            if not gid:
                errors.append(f"RepeatGroups row {rownum}: empty group_id")
                continue
            if gid in groups:
                errors.append(f"RepeatGroups row {rownum}: duplicate group_id '{gid}'")
            sid = str(r.get("section_id", "")).strip()
            if sid not in sections:
                errors.append(f"RepeatGroups row {rownum} ({gid}): unknown section_id '{sid}'")
            try:
                order = float(r.get("order", 0) or 0)
            except (TypeError, ValueError):
                errors.append(f"RepeatGroups row {rownum} ({gid}): 'order' is not a number")
                order = 0
            min_i_raw = r.get("min_instances")
            try:
                min_i = int(min_i_raw) if min_i_raw not in (None, "") else 1
            except (TypeError, ValueError):
                errors.append(f"RepeatGroups row {rownum} ({gid}): min_instances is not an integer")
                min_i = 1
            max_i_raw = r.get("max_instances")
            max_i = None
            if max_i_raw not in (None, ""):
                try:
                    max_i = int(max_i_raw)
                except (TypeError, ValueError):
                    errors.append(f"RepeatGroups row {rownum} ({gid}): max_instances is not an integer")
            if max_i is not None and max_i < min_i:
                errors.append(f"RepeatGroups row {rownum} ({gid}): max_instances < min_instances")
            req_for = norm_list(r.get("required_for_kinds")) if r.get("required_for_kinds") else []
            groups[gid] = {
                "id": gid,
                "section_id": sid,
                "order": order,
                "item_label": str(r.get("item_label", "")).strip() or "Item",
                "min_instances": min_i,
                "max_instances": max_i,
                "add_button_text": str(r.get("add_button_text", "")).strip() or "+ Add another",
                "required_for_kinds": req_for,
                "show_if_raw": (str(r.get("show_if")).strip() if r.get("show_if") else ""),
                "_row": rownum,
                "fields": [],
            }

    # ---- Fields (pass 1: collect; applies_to_kinds/required_for_kinds not yet validated) ----
    fields = {}
    field_order = []
    field_rows_raw = list(sheet_rows(wb["Fields"]))
    for rownum, r in field_rows_raw:
        fid = str(r.get("field_id", "")).strip()
        if not fid:
            errors.append(f"Fields row {rownum}: empty field_id")
            continue
        if fid in fields:
            errors.append(f"Fields row {rownum}: duplicate field_id '{fid}'")
            continue

        ftype = str(r.get("type", "")).strip()
        if ftype not in VALID_TYPES:
            errors.append(f"Fields row {rownum} ({fid}): invalid type '{ftype}' "
                           f"(must be one of {sorted(VALID_TYPES)})")

        sid = str(r.get("section_id", "")).strip()
        if sid not in sections:
            errors.append(f"Fields row {rownum} ({fid}): unknown section_id '{sid}'")

        _gid_raw = r.get("group_id")
        gid = str(_gid_raw).strip() if _gid_raw is not None else ""
        if gid:
            if gid not in groups:
                errors.append(f"Fields row {rownum} ({fid}): unknown group_id '{gid}'")
            elif groups[gid]["section_id"] != sid:
                errors.append(
                    f"Fields row {rownum} ({fid}): section_id '{sid}' does not match "
                    f"group '{gid}''s section_id '{groups[gid]['section_id']}'"
                )

        required = norm_bool(r.get("required"))
        if required not in VALID_REQUIRED:
            errors.append(f"Fields row {rownum} ({fid}): invalid required value '{required}' "
                           f"(must be TRUE, FALSE, or IF_VISIBLE)")

        required_first_only = norm_bool(r.get("required_first_only")) == "TRUE"
        if required_first_only and not gid:
            errors.append(
                f"Fields row {rownum} ({fid}): required_first_only=TRUE has no effect "
                f"without a group_id"
            )

        required_for_kinds = norm_list(r.get("required_for_kinds")) if r.get("required_for_kinds") else []
        if required_for_kinds and required == "TRUE":
            errors.append(
                f"Fields row {rownum} ({fid}): required_for_kinds has no effect when "
                f"required is already TRUE for every kind"
            )

        applies = norm_list(r.get("applies_to_kinds"))

        ui_style = str(r.get("ui_style", "")).strip() if r.get("ui_style") else ""
        if ui_style not in VALID_UI_STYLES:
            errors.append(f"Fields row {rownum} ({fid}): invalid ui_style '{ui_style}' "
                           f"(must be blank or 'toggle')")
        try:
            options = parse_options(r.get("options"))
        except SpecError as e:
            errors.append(f"Fields row {rownum} ({fid}): {e}")
            options = []

        if ftype in OPTION_TYPES and not options:
            errors.append(f"Fields row {rownum} ({fid}): type '{ftype}' requires at least one option")
        if ftype not in OPTION_TYPES and options:
            errors.append(f"Fields row {rownum} ({fid}): type '{ftype}' should not have options")

        try:
            order = float(r.get("order", 0) or 0)
        except (TypeError, ValueError):
            errors.append(f"Fields row {rownum} ({fid}): 'order' is not a number")
            order = 0

        fields[fid] = {
            "id": fid,
            "section_id": sid,
            "group_id": gid or None,
            "order": order,
            "label": str(r.get("label", "")).strip(),
            "type": ftype,
            "options": options,
            "required": required,
            "required_first_only": required_first_only,
            "required_for_kinds": required_for_kinds,
            "applies_to_kinds": applies,
            "ui_style": ui_style,
            "show_if_raw": (str(r.get("show_if")).strip() if r.get("show_if") else ""),
            "disabled_if_raw": (str(r.get("disabled_if")).strip() if r.get("disabled_if") else ""),
            "help_text": str(r.get("help_text", "")).strip() if r.get("help_text") else "",
            "_row": rownum,
        }
        field_order.append(fid)

    # ---- Fields (pass 2: show_if / disabled_if, now that every field_id is known) ----
    for fid, f in fields.items():
        raw = f.pop("show_if_raw")
        f["show_if"] = resolve_show_if(
            raw, fields, errors, where=f"Fields row {f['_row']} ({fid})",
            self_id=fid, self_group_id=f["group_id"],
        )
        disabled_raw = f.pop("disabled_if_raw")
        f["disabled_if"] = resolve_show_if(
            disabled_raw, fields, errors, where=f"Fields row {f['_row']} ({fid}) disabled_if",
            self_id=fid, self_group_id=f["group_id"],
        )

    # ---- RepeatGroups show_if (resolved now that `fields` is fully built) ----
    for gid, g in groups.items():
        raw = g.pop("show_if_raw")
        g["show_if"] = resolve_show_if(
            raw, fields, errors, where=f"RepeatGroups ({gid})",
            self_id=gid, require_groupless_ref=True,
        )

    # ---- KindAxes: field_ids whose OPTION VALUES become valid kind tags ----
    kind_axis_fields = []
    valid_kind_values = set()
    for rownum, r in sheet_rows(wb["KindAxes"]):
        fid = str(r.get("field_id", "")).strip()
        if not fid:
            errors.append(f"KindAxes row {rownum}: empty field_id")
            continue
        if fid not in fields:
            errors.append(f"KindAxes row {rownum}: unknown field_id '{fid}'")
            continue
        if fields[fid]["type"] not in OPTION_TYPES:
            errors.append(
                f"KindAxes row {rownum} ({fid}): type '{fields[fid]['type']}' can't act as a "
                f"kind axis (must be radio/select/checkbox_group)"
            )
            continue
        kind_axis_fields.append(fid)
        valid_kind_values.update(o["value"] for o in fields[fid]["options"])

    def check_kind_refs(values, where):
        for v in values:
            if v != "all" and v not in valid_kind_values:
                errors.append(
                    f"{where}: references unknown kind value '{v}' "
                    f"(not an option of any KindAxes field)"
                )

    # ---- Now validate every applies_to_kinds / required_for_kinds collected earlier ----
    for sid, sec in sections.items():
        check_kind_refs(sec["applies_to_kinds"], f"Sections row {sec['_row']} ({sid})")
    for fid, f in fields.items():
        check_kind_refs(f["applies_to_kinds"], f"Fields row {f['_row']} ({fid})")
        check_kind_refs(f["required_for_kinds"], f"Fields row {f['_row']} ({fid}) required_for_kinds")
    for gid, g in groups.items():
        check_kind_refs(g["required_for_kinds"], f"RepeatGroups ({gid}) required_for_kinds")

    # ---- PDF_Layout (optional overrides; may target a field_id or a group_id) ----
    pdf_overrides = {}
    group_pdf_overrides = {}
    if "PDF_Layout" in wb.sheetnames:
        for rownum, r in sheet_rows(wb["PDF_Layout"]):
            key = str(r.get("field_id_or_group_id", r.get("field_id", ""))).strip()
            entry = {
                "page_break_before": norm_bool(r.get("page_break_before")) == "TRUE",
                "pdf_format": (str(r.get("pdf_format")).strip() if r.get("pdf_format") else None),
            }
            if key in fields:
                pdf_overrides[key] = entry
            elif key in groups:
                group_pdf_overrides[key] = entry
            else:
                errors.append(f"PDF_Layout row {rownum}: '{key}' is not a known field_id or group_id")

    DEFAULT_PDF_FORMAT = {
        "textarea": "paragraph",
        "radio": "checklist",
        "checkbox_group": "checklist",
        "select": "table_row",
        "text": "table_row",
        "email": "table_row",
        "tel": "table_row",
        "date": "table_row",
        "info": "paragraph",
    }

    if errors:
        raise SpecError("\n".join(errors))

    # ---- Assemble ----
    for sec in sections.values():
        sec.pop("_row", None)
    for g in groups.values():
        g.pop("_row", None)

    for fid in field_order:
        f = fields[fid]
        f.pop("_row", None)
        override = pdf_overrides.get(fid, {})
        f["pdf"] = {
            "page_break_before": override.get("page_break_before", False),
            "pdf_format": override.get("pdf_format") or DEFAULT_PDF_FORMAT.get(f["type"], "table_row"),
        }
        if f["group_id"]:
            groups[f["group_id"]]["fields"].append(f)

    # ungrouped fields, bucketed per section
    ungrouped_by_section = {sid: [] for sid in sections}
    for fid in field_order:
        f = fields[fid]
        if not f["group_id"]:
            ungrouped_by_section[f["section_id"]].append(f)

    for gid, g in groups.items():
        g["fields"].sort(key=lambda f: f["order"])
        override = group_pdf_overrides.get(gid, {})
        g["pdf"] = {"page_break_before": override.get("page_break_before", False),
                    "pdf_format": override.get("pdf_format") or "repeated_block"}

    for sid, sec in sections.items():
        items = []
        for f in ungrouped_by_section[sid]:
            items.append({"kind": "field", "order": f["order"], **f})
        for gid, g in groups.items():
            if g["section_id"] == sid:
                items.append({"kind": "group", "order": g["order"], **g})
        items.sort(key=lambda it: it["order"])
        sec["items"] = items
        sec.pop("fields", None)

    ordered_sections = [sections[sid] for sid in sorted(sections, key=lambda s: sections[s]["order"])]

    return {
        "version": 3,
        "kind_axis_fields": kind_axis_fields,
        "sections": ordered_sections,
    }


def main():
    if len(sys.argv) < 2:
        print("usage: python build_schema.py structure.xlsx [output.json]", file=sys.stderr)
        sys.exit(2)
    xlsx_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("form-schema.json")

    try:
        schema = build(xlsx_path)
    except SpecError as e:
        print("Validation failed — nothing was written.\n", file=sys.stderr)
        print(str(e), file=sys.stderr)
        sys.exit(1)

    out_path.write_text(json.dumps(schema, indent=2))

    def count_fields(items):
        n = 0
        for it in items:
            if it["kind"] == "field":
                n += 1
            else:
                n += len(it["fields"])
        return n

    n_fields = sum(count_fields(s["items"]) for s in schema["sections"])
    n_groups = sum(1 for s in schema["sections"] for it in s["items"] if it["kind"] == "group")
    print(f"OK: wrote {out_path}  "
          f"({len(schema['kind_axis_fields'])} kind axes, {len(schema['sections'])} sections, "
          f"{n_fields} fields, {n_groups} repeat group(s))")


if __name__ == "__main__":
    main()
